"""구글 시트(xlsx 내보내기)를 src/data/*.json 으로 변환한다.

사용법:
    pip install openpyxl
    python scripts/import_sheet.py <sheet.xlsx>

시트 xlsx는 구글 시트에서 파일 → 다운로드 → Microsoft Excel(.xlsx)로 받는다.
"""

import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")  # openpyxl DrawingML 경고

OUT_DIR = Path(__file__).resolve().parent.parent / "src" / "data"

EOK = 100_000_000  # 1억 메소


def clean(value):
    """빈 값과 '-' 를 None 으로, 문자열은 양끝 공백 제거."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value in ("", "-", '""'):
            return None
    return value


def to_int(value):
    value = clean(value)
    if isinstance(value, (int, float)):
        return int(value)
    return None


def check(value):
    """'ㅇ'/'O' → 'yes', '△' → 'partial', 'X' → 'no'."""
    value = clean(value)
    if value is None:
        return None
    if value in ("ㅇ", "O", "o"):
        return "yes"
    if value == "△":
        return "partial"
    if value in ("X", "x"):
        return "no"
    return None


def merged_value(ws, row, col):
    """병합된 셀이면 병합 범위의 첫 셀 값을 돌려준다."""
    cell = ws.cell(row, col)
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return ws.cell(rng.min_row, rng.min_col).value
    return cell.value


EQUIP_COLUMNS = {
    "N": "유니온캐 장비 세팅",
    "O": "9보장",
    "P": "얼굴장식",
    "Q": "눈장식",
    "R": "귀고리",
    "S": "반지",
    "T": "뱃지",
    "U": "견장",
    "V": "목걸이",
    "W": "벨트",
    "X": "성배",
    "Y": "3카룻",
    "Z": "네크로",
    "AA": "블랙셋",
    "AB": "앱솔",
}


def import_characters(wb):
    ws = wb["유니온&링크"]
    characters = []
    for row in range(2, ws.max_row + 1):
        get = lambda col: ws[f"{col}{row}"].value  # noqa: E731
        name = clean(get("E"))
        if not name:
            continue
        job_group = clean(merged_value(ws, row, 1))
        note = clean(get("H"))
        equip = {}
        for col, slot in EQUIP_COLUMNS.items():
            text = clean(get(col))
            if text is not None:
                equip[slot] = str(text)
        characters.append(
            {
                "id": name,
                "name": name,
                "jobGroup": job_group,
                "job": clean(get("B")),
                "unionEffect": clean(get("C")),
                "linkEffect": clean(merged_value(ws, row, 4)),
                "plan": clean(get("F")),
                "level": to_int(get("G")),
                "note": note,
                "guild": clean(get("I")),
                "boss": {
                    "karut": check(get("J")),
                    "stoneDesire": check(get("K")),
                    "lucidWillSlime": check(get("L")),
                },
                "unionPower": to_int(get("M")),
                "isMain": note == "본캐",
                # 메이플M 캐릭터는 시트 수식에서도 유니온 레벨에서 빠져 있다.
                "excludeFromUnion": job_group == "특수",
                "equip": equip,
            }
        )
    return characters


def import_target_setups(wb):
    ws = wb["비숍 템세팅"]
    fields = {3: "starforce", 4: "mainPot", 5: "addPot", 6: "flame", 7: "scroll", 8: "budget"}
    slots = []
    for col in range(2, ws.max_column + 1):
        # 시드링 칸처럼 부위 행이 비어 있으면 세트 헤더를 부위 이름으로 쓴다
        slot = clean(ws.cell(2, col).value) or clean(ws.cell(1, col).value)
        if not slot:
            continue
        # 세트 이름이 첫 부위 칸에만 적혀 있으면 앞 부위의 세트를 이어받는다
        group = clean(merged_value(ws, 1, col)) or (slots[-1]["group"] if slots else None)
        entry = {"group": group, "slot": slot}
        for row, key in fields.items():
            value = clean(ws.cell(row, col).value)
            if isinstance(value, float):  # 작 15% 가 0.15 로 저장돼 있음
                value = f"{round(value * 100)}%"
            entry[key] = value
        notes = [clean(ws.cell(r, col).value) for r in range(9, 12)]
        entry["note"] = " / ".join(n for n in notes if n) or None
        slots.append(entry)
    advice = [clean(ws.cell(r, 3).value) for r in range(47, 54)]
    return [
        {
            "characterId": "비숍자그",
            "title": "비숍 템세팅",
            "slots": slots,
            "note": "\n".join(a for a in advice if a),
        }
    ]


def parse_date(text):
    """'2023.09.02', '23.09.14', '24.04.14(?)' → 'YYYY-MM-DD'. 일자를 모르면('23.10.??') 'YYYY-MM'."""
    m = re.match(r"(\d{2,4})\.(\d{1,2})(?:\.(\d{1,2}))?", text)
    if not m:
        return None
    year, month, day = m.groups()
    if len(year) == 2:
        year = "20" + year
    if day is None:
        return f"{year}-{int(month):02d}"
    return f"{year}-{int(month):02d}-{int(day):02d}"


def parse_eok(text):
    """'155억', '199.99억', '675억+10%관세' → 메소 정수 (첫 숫자 기준)."""
    m = re.search(r"([\d.]+)\s*억", text)
    if not m:
        return None
    try:
        return round(float(m.group(1).rstrip(".")) * EOK)
    except ValueError:
        return None


def import_purchases(wb):
    ws = wb["영수증..."]
    purchases = []
    for row in range(2, ws.max_row + 1):
        date_text = clean(ws.cell(row, 1).value)
        item = clean(ws.cell(row, 2).value)
        if not date_text or not item:
            continue
        date_text = str(date_text)
        price_text = str(clean(ws.cell(row, 4).value) or "")
        purchases.append(
            {
                "id": f"p{len(purchases) + 1}",
                "date": parse_date(date_text),
                "dateText": date_text,
                "itemName": item,
                "detail": clean(ws.cell(row, 3).value),
                "price": parse_eok(price_text),
                "priceText": price_text or None,
                "comment": clean(ws.cell(row, 5).value),
                # 어느 캐릭터 구매인지는 시트에 없어서 비워 둔다.
                "characterId": None,
            }
        )
    return purchases


# 메소 가격이 아니라 1억 메소당 원화 시세인 품목
WON_RATE_ITEMS = ("물통", "메소 마켓")


def import_prices(wb):
    ws = wb["mvp작 참고용 옥션 시세표"]
    dates = {}
    for col in range(3, ws.max_column + 1):
        header = clean(ws.cell(1, col).value)
        if header and re.match(r"\d{4}\.\d{2}\.\d{2}", str(header)):
            dates[col] = str(header).replace(".", "-")
    quotes = []
    item = None
    for row in range(2, ws.max_row + 1):
        name = clean(ws.cell(row, 1).value)
        bundle = to_int(ws.cell(row, 2).value)
        if name:
            item = name
        if not item or not bundle:
            continue
        for col, date in dates.items():
            value = ws.cell(row, col).value
            if not isinstance(value, (int, float)):
                continue
            if item in WON_RATE_ITEMS:
                price, unit = round(value), "won_per_eok"
            else:
                # 시트 값은 세트 전체 가격(억)
                price, unit = round(value * EOK), "meso"
            quotes.append(
                {"item": item, "bundleSize": bundle, "date": date, "price": price, "unit": unit}
            )
    return quotes


def import_todos():
    # 강화 탭은 메모 몇 줄이라 직접 옮긴다.
    items = [
        ("그리드 펜던트 드랍템으로 올리기", None),
        ("데벤져 가성비 템 세팅 해주기", None),
        ("바이퍼 상의 수정해주기", "자그3 인벤토리에 넣어둠"),
        ("주흔 작", "12/10 토드"),
    ]
    return [
        {"id": f"t{i + 1}", "text": text, "note": note, "done": False}
        for i, (text, note) in enumerate(items)
    ]


def keep_links(purchases):
    """웹에서 연결해 둔 구매 ↔ 캐릭터 연결은 다시 가져와도 유지한다."""
    path = OUT_DIR / "purchases.json"
    if not path.exists():
        return purchases
    old = {(p["dateText"], p["itemName"], p["detail"]): p.get("characterId") for p in json.loads(path.read_text("utf-8"))}
    for p in purchases:
        p["characterId"] = old.get((p["dateText"], p["itemName"], p["detail"]))
    return purchases


def write(name, data):
    path = OUT_DIR / f"{name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{path.relative_to(OUT_DIR.parent.parent)}: {len(data)}건")


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    # data_only: 수식 대신 마지막으로 계산된 값을 읽는다
    wb = openpyxl.load_workbook(src, data_only=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write("characters", import_characters(wb))
    write("targetSetups", import_target_setups(wb))
    write("purchases", keep_links(import_purchases(wb)))
    write("prices", import_prices(wb))
    write("todos", import_todos())


if __name__ == "__main__":
    main()
